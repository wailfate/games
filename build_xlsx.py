"""games.json -> 게임_홍보링크.xlsx (게임이 늘면 games.json 만 고치고 다시 돌린다)."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
SITE = "https://wailfate.github.io/games"
games = json.loads((ROOT / "games.json").read_text(encoding="utf-8"))["games"]

wb = Workbook()
ws = wb.active
ws.title = "게임 홍보 링크"
head = ["No", "게임 이름", "장르", "한 줄 소개", "패키지명", "플레이 스토어 주소",
        "홍보 페이지 주소", "카톡 복붙 문구 (?src=kakao)", "비고"]
ws.append(head)
for i, g in enumerate(games, 1):
    promo = f"{SITE}/{g['id']}/"
    ws.append([
        i, g["name"], g["genre"], g["tagline"], g["package"],
        f"https://play.google.com/store/apps/details?id={g['package']}",
        promo,
        f"'{g['name']}'에 대해 알아보세요 - {promo}?src=kakao",
        "",
    ])
ws.append([])
ws.append(["", "게임 상점", "", "", "", "", f"{SITE}/", "", ""])

fill = PatternFill("solid", fgColor="1F1A13")
for c in ws[1]:
    c.font = Font(bold=True, color="F3E3B5")
    c.fill = fill
    c.alignment = Alignment(horizontal="center", vertical="center")
widths = [5, 30, 16, 60, 28, 62, 46, 70, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
for r in range(2, ws.max_row + 1):
    for col in (6, 7):
        cell = ws.cell(r, col)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")
ws.freeze_panes = "C2"

out = ROOT / "게임_홍보링크.xlsx"
wb.save(out)
print("saved", out, len(games), "games")
